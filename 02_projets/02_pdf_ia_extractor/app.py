import io
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from pypdf import PdfReader
import streamlit as st

st.set_page_config(page_title="PDF Extractor Pro - Enterprise", layout="wide")

st.title("📄 Saisie Automatique & Synthèse Comptable PDF")
st.subheader("Numérisez vos factures, consolidez vos dépenses et générez un rapport financier parfait.")


def extract_data_from_pdf(pdf_file):
    reader = PdfReader(pdf_file)
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""

    num_facture = re.search(r"FACTURE N°\s*:\s*([A-Z0-9-]+)", text)
    fournisseur = re.search(r"Fournisseur\s*:\s*(.*)", text)
    siret = re.search(r"SIRET\s*:\s*([\d\s]+)", text)
    ville = re.search(r"Ville\s*:\s*(.*)", text)
    date_fac = re.search(r"Date\s*:\s*([\d/]+)", text)
    echeance = re.search(r"Échéance\s*:\s*([\d/]+)", text)
    client = re.search(r"Facturé à\s*:\s*(.*)", text)
    categorie = re.search(r"Catégorie\s*:\s*(.*)", text)
    total_ht = re.search(r"Total HT\s*:\s*([\d\.]+)", text)
    tva = re.search(r"TVA \(20%\)\s*:\s*([\d\.]+)", text)
    total_ttc = re.search(r"Total TTC\s*:\s*([\d\.]+)", text)
    statut = re.search(r"Statut\s*:\s*([^|]+)", text)
    reglement = re.search(r"Mode de Règlement\s*:\s*(.*)", text)

    return {
        "N° Facture": num_facture.group(1).strip() if num_facture else "N/A",
        "Date": date_fac.group(1).strip() if date_fac else "N/A",
        "Échéance": echeance.group(1).strip() if echeance else "N/A",
        "Fournisseur": fournisseur.group(1).strip() if fournisseur else "N/A",
        "SIRET": siret.group(1).strip() if siret else "N/A",
        "Ville": ville.group(1).strip() if ville else "N/A",
        "Client": client.group(1).strip() if client else "N/A",
        "Catégorie": categorie.group(1).strip() if categorie else "N/A",
        "Montant HT (€)": float(total_ht.group(1)) if total_ht else 0.0,
        "TVA (€)": float(tva.group(1)) if tva else 0.0,
        "Montant TTC (€)": float(total_ttc.group(1)) if total_ttc else 0.0,
        "Statut": statut.group(1).strip() if statut else "N/A",
        "Règlement": reglement.group(1).strip() if reglement else "N/A",
        "Fichier": pdf_file.name,
        "Texte Brut": text,
    }


def generate_styled_excel(df):
    output = io.BytesIO()
    df_export = df.drop(columns=["Texte Brut", "Fichier"])

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_export.to_excel(writer, index=False, sheet_name="Journal des Achats")
        worksheet = writer.sheets["Journal des Achats"]

        # Styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        paye_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        paye_font = Font(name="Calibri", size=10, color="375623", bold=True)
        
        attente_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        attente_font = Font(name="Calibri", size=10, color="7F6000", bold=True)

        a_payer_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        a_payer_font = Font(name="Calibri", size=10, color="C65911", bold=True)

        center_align = Alignment(horizontal="center", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        worksheet.views.sheetView[0].showGridLines = True
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        # En-tête
        for col_num in range(1, len(df_export.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Lignes de données
        for row_num in range(2, len(df_export) + 2):
            for col_num, col_name in enumerate(df_export.columns, 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.border = border_thin
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = center_align

                if " (€)" in col_name:
                    cell.number_format = '#,##0.00 "€"'
                    cell.alignment = right_align

                if col_name == "Statut":
                    val = str(cell.value).strip()
                    if val == "Payé":
                        cell.fill = paye_fill
                        cell.font = paye_font
                    elif val == "En attente":
                        cell.fill = attente_fill
                        cell.font = attente_font
                    else:
                        cell.fill = a_payer_fill
                        cell.font = a_payer_font

        # Ligne de TOTAL au bas du tableau
        total_row = len(df_export) + 2
        worksheet.cell(row=total_row, column=1, value="TOTAL GENERAL").font = Font(name="Calibri", size=11, bold=True)
        
        # Formules de somme
        ht_col = df_export.columns.get_loc("Montant HT (€)") + 1
        tva_col = df_export.columns.get_loc("TVA (€)") + 1
        ttc_col = df_export.columns.get_loc("Montant TTC (€)") + 1

        for c_idx in [ht_col, tva_col, ttc_col]:
            col_letter = get_column_letter(c_idx)
            cell = worksheet.cell(row=total_row, column=c_idx, value=f"=SUM({col_letter}2:{col_letter}{total_row-1})")
            cell.font = Font(name="Calibri", size=11, bold=True)
            cell.number_format = '#,##0.00 "€"'
            cell.border = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

        # Largeurs de colonnes
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = max(max_len + 5, 14)

    return output.getvalue()


uploaded_files = st.file_uploader("Glissez vos factures PDF ici", type=["pdf"], accept_multiple_files=True)

if uploaded_files:
    data_list = [extract_data_from_pdf(pdf) for pdf in uploaded_files]
    df = pd.DataFrame(data_list)

    # Indicateurs Métriques
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Factures traitées", len(df))
    m2.metric("Total HT", f"{df['Montant HT (€)'].sum():,.2f} €")
    m3.metric("Total TVA", f"{df['TVA (€)'].sum():,.2f} €")
    m4.metric("Total TTC", f"{df['Montant TTC (€)'].sum():,.2f} €")

    st.markdown("---")

    st.success(" Extraction et analyse financière réussies !")
    st.dataframe(df.drop(columns=["Texte Brut"]), use_container_width=True)

    # Inspecteur
    st.markdown("---")
    col_insp1, col_insp2 = st.columns(2)
    
    with col_insp1:
        st.subheader("🔍 Inspecteur de Document")
        selected_file = st.selectbox("Sélectionnez une facture à vérifier :", df["Fichier"].tolist())
        row = df[df["Fichier"] == selected_file].iloc[0]
        st.json({
            "Fournisseur": row["Fournisseur"],
            "SIRET": row["SIRET"],
            "Client": row["Client"],
            "Catégorie": row["Catégorie"],
            "Montant TTC": f"{row['Montant TTC (€)']} €",
            "Statut": row["Statut"]
        })

    with col_insp2:
        st.subheader("📊 Répartition des Dépenses par Catégorie")
        st.bar_chart(df.groupby("Catégorie")["Montant TTC (€)"].sum())

    # Export
    excel_data = generate_styled_excel(df)
    st.markdown("---")
    st.download_button(
        label="📥 Télécharger le Journal d'Achats Complexe (.xlsx)",
        data=excel_data,
        file_name="Journal_Achats_Entreprise.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )