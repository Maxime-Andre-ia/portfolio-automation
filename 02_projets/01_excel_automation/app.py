import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="AutoExcel - Nettoyage & Consolidation", layout="wide"
)

st.title("📊 Automatisation & Consolidation Excel")
st.subheader("Transformez vos fichiers dispersés en un rapport propre en 1 clic.")


def generate_styled_excel(df):
    """Génère un fichier Excel parfaitement mis en forme avec filtres et centrage."""
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Rapport Consolidé")
        workbook = writer.book
        worksheet = writer.sheets["Rapport Consolidé"]

        # --- STYLES ET COULEURS ---
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        paye_fill = PatternFill(
            start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
        )
        paye_font = Font(name="Calibri", size=10, color="375623", bold=True)

        attente_fill = PatternFill(
            start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
        )
        attente_font = Font(name="Calibri", size=10, color="7F6000", bold=True)

        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        center_alignment = Alignment(horizontal="center", vertical="center")

        # Activer le quadrillage, le filtre automatique (menu déroulant) et figer la 1ère ligne
        worksheet.views.sheetView[0].showGridLines = True
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions  # Activer les flèches de filtre

        # 1. En-tête centré avec style
        for col_num, col_name in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment

        # 2. Données entièrement centrées
        for row_num in range(2, len(df) + 2):
            for col_num, col_name in enumerate(df.columns, 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.border = border_thin
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = center_alignment  # Centrage absolu pour toutes les colonnes

                # Formatage monétaire
                if col_name == "Montant HT":
                    cell.number_format = '#,##0.00 "€"'

                # Couleurs conditionnelles sur le statut
                if col_name == "Statut":
                    val_statut = str(cell.value).strip()
                    if val_statut == "Payé":
                        cell.fill = paye_fill
                        cell.font = paye_font
                    else:
                        cell.fill = attente_fill
                        cell.font = attente_font

        # 3. Ajustement automatique de la largeur des colonnes
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    val_str = str(cell.value)
                    if cell.number_format == '#,##0.00 "€"' and isinstance(
                        cell.value, (int, float)
                    ):
                        val_str = f"{cell.value:,.2f} €"
                    max_len = max(max_len, len(val_str))
            # On ajoute un peu de marge pour laisser la place aux flèches de filtre
            worksheet.column_dimensions[col_letter].width = max(max_len + 6, 14)

    return output.getvalue()


# --- INTERFACE STREAMLIT ---
uploaded_files = st.file_uploader(
    "Déposez vos fichiers Excel de la semaine ici",
    type=["xlsx"],
    accept_multiple_files=True,
)

if uploaded_files:
    df_list = [pd.read_excel(f) for f in uploaded_files]
    raw_df = pd.concat(df_list, ignore_index=True)

    st.info(
        f"📁 {len(uploaded_files)} fichiers chargés ({len(raw_df)} lignes"
        " brutes)."
    )

    if st.button("⚡ Lancer le traitement et le nettoyage"):
        df_clean = raw_df.copy()

        # Nettoyage
        df_clean["Statut"] = (
            df_clean["Statut"].astype(str).str.strip().str.capitalize()
        )
        df_clean["Statut"] = df_clean["Statut"].replace({"Paye": "Payé"})
        df_clean = df_clean.drop_duplicates()
        df_clean["Montant HT"] = df_clean["Montant HT"].fillna(0)
        df_clean["Date"] = pd.to_datetime(
            df_clean["Date"], format="mixed"
        ).dt.strftime("%d/%m/%Y")

        # Indicateurs
        col1, col2, col3 = st.columns(3)
        col1.metric(
            "Lignes nettoyées",
            len(df_clean),
            delta=f"-{len(raw_df) - len(df_clean)} doublons",
        )
        col2.metric(
            "Chiffre d'Affaires HT", f"{df_clean['Montant HT'].sum():,.2f} €"
        )
        col3.metric(
            "Factures réglées", len(df_clean[df_clean["Statut"] == "Payé"])
        )

        st.success("Traitement et mise en forme terminés !")
        st.dataframe(df_clean, use_container_width=True)

        # Génération du fichier Excel stylisé
        excel_data = generate_styled_excel(df_clean)

        st.download_button(
            label="📥 Télécharger le rapport consolidé design (.xlsx)",
            data=excel_data,
            file_name="Rapport_Consolide_Design.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )